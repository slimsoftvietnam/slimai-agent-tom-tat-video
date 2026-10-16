"""Bước 5: dựng tài liệu Word (+ PDF nếu máy có Pages hoặc LibreOffice) từ noi_dung.json, rồi tự chấm.

Cách dùng:
  python dung_tai_lieu.py <thư-mục-kết-quả> [--nguoi-soan "Tên bạn"] [--khong-pdf]

Đọc: noi_dung.json (agent viết), loi_thoai.txt, thong_tin.json, video.mp4
Tạo: <tên>.docx, <tên>.pdf (nếu được), hinh/*.jpg|png (ảnh chụp + sơ đồ + bìa), tu_cham.txt
Phong cách "Ghi chú khoá học": bìa nền tối nhấn cam, chữ thân có chân, tiêu đề chương cam, hộp nền kem.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import platform
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ve_minh_hoa  # noqa: E402

ORANGE = RGBColor(0xE8, 0x69, 0x2C)
INK = RGBColor(0x22, 0x20, 0x1E)
BODY = RGBColor(0x2E, 0x2B, 0x28)
MUTED = RGBColor(0x7A, 0x74, 0x6B)
CREAM = "F7F1E6"
SANS = "Arial"
SERIF = "Georgia"
PAGE_W, PAGE_H = Inches(8.27), Inches(11.69)  # A4


def secs(ts) -> int | None:
    if not isinstance(ts, str) or not re.fullmatch(r"(\d{1,2}:)?\d{1,2}:\d{2}", ts.strip()):
        return None
    p = [int(x) for x in ts.strip().split(":")]
    return p[0] * 3600 + p[1] * 60 + p[2] if len(p) == 3 else p[0] * 60 + p[1]


def stamp(s: int) -> str:
    h, r = divmod(int(s), 3600)
    m, s = divmod(r, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def slug(text: str) -> str:
    import unicodedata
    text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    text = "".join(c for c in text if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")[:60] or "tai-lieu"


# ---------- định dạng ----------

def run(par, text, size=11, bold=False, italic=False, color=BODY, font=SERIF, spacing=None):
    r = par.add_run(text)
    r.font.name = font
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    r.font.color.rgb = color
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        r._element.get_or_add_rPr().append(sp)
    return r


def link(par, url, text, size=8.5):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), SANS)
    fonts.set(qn("w:hAnsi"), SANS)
    rpr.append(fonts)
    for tag, val in (("w:b", None), ("w:color", "E8692C"), ("w:sz", str(int(size * 2)))):
        el = OxmlElement(tag)
        if val:
            el.set(qn("w:val"), val)
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def shade(par, fill=CREAM):
    shd = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", fill)):
        shd.set(qn(f"w:{k}"), v)
    par._p.get_or_add_pPr().append(shd)


def border(par, side="bottom", color="E8692C", size="8", space="4"):
    ppr = par._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    b = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", size), ("space", space), ("color", color)):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)


def field(par, code, size=8):
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), code)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:b", None), ("w:color", "E8692C"), ("w:sz", str(size * 2))):
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), SANS)
            el.set(qn("w:hAnsi"), SANS)
        elif val:
            el.set(qn("w:val"), val)
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    f.append(r)
    par._p.append(f)


def para(doc, text="", size=11, italic=False, color=BODY, after=7, align=None, font=SERIF, bold=False):
    p = doc.add_paragraph()
    if text:
        run(p, text, size=size, italic=italic, color=color, font=font, bold=bold)
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing = 1.38
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, level, moc=None, page_break=False):
    p = doc.add_heading(level=level)
    if level == 1:
        run(p, text, size=19, bold=True, color=ORANGE, font=SANS)
        border(p, color="E6DCCB", size="6", space="6")
    else:
        run(p, text, size=13, bold=True, color=INK, font=SANS)
    if moc:
        run(p, f"  ·  {moc}", size=10 if level > 1 else 12, bold=True, color=MUTED, font=SANS)
    pf = p.paragraph_format
    pf.space_before = Pt(18 if level == 1 else 13)
    pf.space_after = Pt(7)
    pf.keep_with_next = True
    pf.page_break_before = page_break
    return p


def caption(doc, label, moc, title, desc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, label, size=8, bold=True, color=ORANGE, font=SANS, spacing=30)
    if moc:
        run(p, f"  ·  {moc}", size=8, bold=True, color=ORANGE, font=SANS)
    if title:
        run(p, f"  —  {title}", size=8.5, bold=True, color=INK, font=SANS)
    p.paragraph_format.space_after = Pt(0)
    if desc:
        q = para(doc, desc, size=9.5, italic=True, color=MUTED, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        q.paragraph_format.line_spacing = 1.2
    else:
        p.paragraph_format.space_after = Pt(12)


def picture(doc, path: Path, width_in=6.0):
    doc.add_picture(str(path), width=Inches(width_in))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True


def full_page_image(doc, path: Path):
    """Ảnh phủ kín trang đầu (neo theo trang, nằm sau chữ) — cách duy nhất Pages hiển thị đúng bìa tràn lề."""
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    r = p.add_run()
    r.add_picture(str(path), width=Emu(PAGE_W), height=Emu(PAGE_H))
    inline = r._element.find(".//" + qn("wp:inline"))
    graphic = copy.deepcopy(inline.find(qn("a:graphic")))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int(PAGE_W)}" cy="{int(PAGE_H)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="9001" name="Bia"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)
    r.add_break(WD_BREAK.PAGE)


# ---------- dữ liệu ----------

def video_link(base: str, s: int) -> str | None:
    return f"{base}{'&' if '?' in base else '?'}t={s}s" if base else None


def grab(video: Path, s: int, dest: Path) -> bool:
    subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-ss", str(s), "-i", str(video),
                    "-frames:v", "1", "-vf", "scale=1280:-2", "-q:v", "3", str(dest)], check=False)
    return dest.exists() and dest.stat().st_size > 0


def to_pdf(docx: Path) -> Path | None:
    pdf = docx.with_suffix(".pdf")
    pdf.unlink(missing_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       capture_output=True, timeout=300)
    elif platform.system() == "Darwin" and Path("/Applications/Pages.app").exists():
        script = (f'tell application "Pages"\nset d to open POSIX file "{docx}"\ndelay 2\n'
                  f'export d to POSIX file "{pdf}" as PDF\nclose d saving no\nend tell')
        subprocess.run(["osascript", "-e", script], capture_output=True, timeout=300)
    return pdf if pdf.exists() else None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("thu_muc")
    ap.add_argument("--nguoi-soan", default="")
    ap.add_argument("--khong-pdf", action="store_true")
    a = ap.parse_args()
    out = Path(a.thu_muc).resolve()
    nd = json.loads((out / "noi_dung.json").read_text(encoding="utf-8"))
    info = json.loads((out / "thong_tin.json").read_text(encoding="utf-8"))
    transcript = (out / "loi_thoai.txt").read_text(encoding="utf-8")
    valid = {secs(t) for t in re.findall(r"^\[((?:\d{2}:)?\d{2}:\d{2})\]", transcript, re.M)}
    duration = info.get("thoi_luong_giay") or max(valid or {0})
    base = info.get("link", "")
    video = out / "video.mp4"
    title = nd.get("tieu_de") or info.get("tieu_de", "Tài liệu từ video")
    nd.setdefault("nam", str(dt.date.today().year))
    chapters = nd.get("chuong", [])
    questions = nd.get("cau_hoi", [])
    terms = nd.get("thuat_ngu", [])
    warnings: list[str] = []
    has_browser = ve_minh_hoa.find_browser() is not None
    if not has_browser:
        warnings.append("Không tìm thấy Chrome/Edge: không vẽ được bìa và sơ đồ (cài Google Chrome rồi chạy lại)")

    img_dir = out / "hinh"
    img_dir.mkdir(exist_ok=True)
    for old in list(img_dir.glob("*.jpg")) + list(img_dir.glob("*.png")):
        old.unlink()

    # Đếm trước số hình để in lên bìa
    planned = sum(len(m.get("hinh", [])) + len(m.get("minh_hoa", [])) for c in chapters for m in c.get("muc", []))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.left_margin = sec.right_margin = Inches(0.95)
    sec.top_margin, sec.bottom_margin = Inches(0.95), Inches(0.85)
    sec.different_first_page_header_footer = True

    # Bìa
    cover_shot = None
    cm = secs((nd.get("bia") or {}).get("hinh_moc", ""))
    if cm is not None and cm in valid and video.exists() and grab(video, cm, img_dir / "bia-anh.jpg"):
        cover_shot = img_dir / "bia-anh.jpg"
    cover_png = img_dir / "bia.png"
    if has_browser and ve_minh_hoa.render_cover(nd, info, planned, len(questions), cover_shot, cover_png):
        full_page_image(doc, cover_png)
    else:
        for _ in range(6):
            para(doc)
        para(doc, "GHI CHÚ KHOÁ HỌC", size=10, bold=True, color=ORANGE, font=SANS)
        para(doc, title, size=30, bold=True, color=INK, font=SANS, after=10)
        if nd.get("phu_de"):
            para(doc, nd["phu_de"], size=13, italic=True, color=MUTED)
        border(para(doc), size="18")
        para(doc, f"Nguồn: {info.get('tieu_de', '')}" + (f" — {info['kenh']}" if info.get("kenh") else ""), size=10, color=MUTED, font=SANS)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Đầu trang, chân trang (trang bìa không có)
    hp = sec.header.paragraphs[0]
    run(hp, "GHI CHÚ KHOÁ HỌC  /  ", size=7.5, bold=True, color=MUTED, font=SANS, spacing=20)
    short = title if len(title) <= 55 else title[:52].rsplit(" ", 1)[0].rstrip(":,;–—- ") + "…"
    run(hp, short.upper(), size=7.5, bold=True, color=MUTED, font=SANS, spacing=20)
    border(hp, color="E8692C", size="6", space="4")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "TÀI LIỆU HỌC TẬP  ·  ", size=7.5, bold=True, color=MUTED, font=SANS, spacing=20)
    field(fp, "PAGE")

    # Mục lục
    heading(doc, "Mục lục", 1)
    for i, ch in enumerate(chapters, 1):
        p = para(doc, after=2)
        run(p, f"{i:02d}   ", size=11, bold=True, color=ORANGE, font=SANS)
        run(p, ch.get("tieu_de", ""), size=11, bold=True, color=INK, font=SANS)
        if ch.get("moc"):
            run(p, f"   {ch['moc']}", size=9, color=MUTED, font=SANS)
        q = p
        for j, m in enumerate(ch.get("muc", []), 1):
            q = para(doc, after=0)
            q.paragraph_format.left_indent = Inches(0.42)
            run(q, f"{i}.{j}  {m.get('tieu_de', '')}", size=9.5, color=MUTED, font=SANS)
        q.paragraph_format.space_after = Pt(6)
    for name in ("Thuật ngữ cần nhớ", "Điểm cần ghi nhớ", "Câu hỏi tự ôn"):
        p = para(doc, after=2)
        run(p, name, size=11, bold=True, color=INK, font=SANS)

    # Tổng quan
    heading(doc, "Tổng quan", 1, page_break=True)
    for block in re.split(r"\n\s*\n", nd.get("tong_quan", "").strip()):
        if block.strip():
            p = para(doc, block.strip(), after=0)
            shade(p)
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.05)
            border(p, side="left", color="E8692C", size="18", space="8")
    para(doc, after=4)
    if nd.get("y_chinh"):
        heading(doc, "Ý chính", 2)
        for y in nd["y_chinh"]:
            p = doc.add_paragraph(style="List Bullet")
            run(p, y)
            p.paragraph_format.space_after = Pt(4)

    # Nội dung
    words = images = diagrams = 0
    last_moc = 0
    sections_without_moc = 0
    chapters_without_visual = []
    for i, ch in enumerate(chapters, 1):
        heading(doc, f"{i:02d}  {ch.get('tieu_de', '')}", 1, moc=ch.get("moc"), page_break=True)
        if ch.get("mo_dau"):
            para(doc, ch["mo_dau"], size=11.5, italic=True, color=MUTED, after=10)
        visuals_here = 0
        for j, m in enumerate(ch.get("muc", []), 1):
            s = secs(m.get("moc", ""))
            heading(doc, f"{i}.{j}  {m.get('tieu_de', '')}", 2, moc=stamp(s) if s is not None else None)
            if s is None:
                sections_without_moc += 1
                warnings.append(f"Mục {i}.{j} thiếu mốc thời gian hợp lệ")
            else:
                last_moc = max(last_moc, s)
                url = video_link(base, s)
                if url:
                    p = para(doc, after=4)
                    run(p, "▶  ", size=8, color=ORANGE, font=SANS)
                    link(p, url, f"Xem đoạn này trong video ({stamp(s)})")
            # Hình chụp màn hình: chèn sau đoạn văn có mốc gần nhất
            pending = []
            for h in m.get("hinh", []):
                hs = secs(h.get("moc", ""))
                if hs is None or (valid and hs not in valid) or hs >= max(duration, 1):
                    warnings.append(f"Bỏ hình mục {i}.{j}: mốc '{h.get('moc')}' không có trong lời thoại")
                    continue
                pending.append((hs, h))
            pending.sort(key=lambda x: x[0])
            # Sơ đồ tự vẽ: chèn sau đoạn thứ "sau_doan" (mặc định sau đoạn đầu)
            drawn = {}
            for k, mh in enumerate(m.get("minh_hoa", [])):
                drawn.setdefault(int(mh.get("sau_doan", 1)), []).append(mh)
            doans = m.get("doan", [])
            for k, d in enumerate(doans, 1):
                text = d.get("noi_dung", "") if isinstance(d, dict) else str(d)
                para(doc, text)
                words += len(text.split())
                if isinstance(d, dict) and secs(d.get("moc", "")) is not None:
                    last_moc = max(last_moc, secs(d["moc"]))
                nxt = secs(doans[k].get("moc", "")) if k < len(doans) and isinstance(doans[k], dict) else None
                while pending and (nxt is None or pending[0][0] < nxt):
                    hs, h = pending.pop(0)
                    img = img_dir / f"hinh-{images + diagrams + 1:02d}-{stamp(hs).replace(':', '')}.jpg"
                    if not grab(video, hs, img):
                        warnings.append(f"Không chụp được hình tại {stamp(hs)}")
                        continue
                    images += 1
                    visuals_here += 1
                    picture(doc, img)
                    caption(doc, f"HÌNH {images + diagrams:02d}", stamp(hs), h.get("tieu_de", ""), h.get("chu_thich", ""))
                for mh in drawn.pop(k, []) + (drawn.pop(len(doans), []) if k == len(doans) else []):
                    png = img_dir / f"so-do-{images + diagrams + 1:02d}.png"
                    if has_browser and ve_minh_hoa.render_diagram(mh, png):
                        diagrams += 1
                        visuals_here += 1
                        picture(doc, png, 6.2)
                        caption(doc, f"SƠ ĐỒ {images + diagrams:02d}", None, mh.get("tieu_de", ""), mh.get("chu_thich", ""))
                    else:
                        warnings.append(f"Không vẽ được sơ đồ '{mh.get('tieu_de')}' (loại '{mh.get('loai')}') ở mục {i}.{j}")
            for left in drawn.values():
                for mh in left:
                    warnings.append(f"Sơ đồ '{mh.get('tieu_de')}' ở mục {i}.{j}: sau_doan vượt quá số đoạn, bị bỏ")
        if ch.get("tom_tat"):
            p = para(doc, after=0)
            run(p, "TÓM TẮT CHƯƠNG", size=8.5, bold=True, color=ORANGE, font=SANS, spacing=30)
            shade(p)
            p.paragraph_format.space_before = Pt(10)
            for t in ch["tom_tat"]:
                q = para(doc, after=0, size=10.5)
                run(q, "—  ", size=10.5, bold=True, color=ORANGE, font=SANS)
                run(q, t, size=10.5)
                shade(q)
            para(doc, after=6)
        if not visuals_here:
            chapters_without_visual.append(i)

    # Phụ lục
    if terms:
        heading(doc, "Thuật ngữ cần nhớ", 1, page_break=True)
        for t in terms:
            p = doc.add_paragraph(style="List Bullet")
            run(p, t.get("tu", ""), bold=True, color=INK)
            run(p, f" — {t.get('giai_thich', '')}")
            p.paragraph_format.space_after = Pt(4)
    if nd.get("ghi_nho"):
        heading(doc, "Điểm cần ghi nhớ", 1)
        for g in nd["ghi_nho"]:
            p = doc.add_paragraph(style="List Bullet")
            run(p, g)
            p.paragraph_format.space_after = Pt(4)
    if questions:
        heading(doc, "Câu hỏi tự ôn", 1)
        for q_ in questions:
            p = doc.add_paragraph(style="List Number")
            run(p, q_)
            p.paragraph_format.space_after = Pt(4)
    src = para(doc, after=0, size=9)
    src.paragraph_format.space_before = Pt(16)
    run(src, "Nguồn: ", size=9, bold=True, color=INK, font=SANS)
    run(src, f"{info.get('tieu_de', '')}" + (f" — {info['kenh']}" if info.get("kenh") else ""), size=9, color=MUTED, font=SANS)
    if base:
        run(src, "  ", size=9)
        link(src, base, base, size=8.5)
    if a.nguoi_soan:
        para(doc, f"Biên soạn: {a.nguoi_soan} · {dt.date.today():%d/%m/%Y}", size=9, color=MUTED, font=SANS)

    name = slug(title)
    docx_path = out / f"{name}.docx"
    doc.save(docx_path)
    pdf_path = None if a.khong_pdf else to_pdf(docx_path)

    minutes = max(duration / 60, 1)
    sections = sum(len(c.get("muc", [])) for c in chapters)
    visuals = images + diagrams
    checks = [
        ("Đủ nội dung (≥ 70 chữ / phút video)", words / minutes >= 70, f"{words} chữ, {words / minutes:.0f} chữ/phút"),
        ("Phủ hết video (mục cuối ≥ 85% thời lượng)", last_moc >= 0.85 * duration - 60, f"mốc cuối {stamp(last_moc)} / {stamp(duration)}"),
        ("Mỗi mục có mốc thời gian", sections_without_moc == 0, f"{sections - sections_without_moc}/{sections} mục"),
        ("Có hình minh hoạ (≥ 1 hình / 5 phút)", visuals >= max(1, int(minutes // 5)), f"{images} ảnh chụp + {diagrams} sơ đồ"),
        ("Chương nào cũng có hình minh hoạ", not chapters_without_visual,
         "đủ" if not chapters_without_visual else "thiếu ở chương " + ", ".join(map(str, chapters_without_visual))),
        ("Có thuật ngữ (≥ 5)", len(terms) >= 5, f"{len(terms)} thuật ngữ"),
        ("Có câu hỏi ôn tập (≥ 5)", len(questions) >= 5, f"{len(questions)} câu"),
    ]
    report = ["BẢNG TỰ CHẤM", *(f"[{'ĐẠT' if ok else 'CHƯA ĐẠT'}] {n}: {d}" for n, ok, d in checks)]
    if warnings:
        report += ["", "Cảnh báo:", *("- " + w for w in warnings)]
    (out / "tu_cham.txt").write_text("\n".join(report), encoding="utf-8")
    print("\n".join(report))
    print(f"\nWORD: {docx_path}")
    print(f"PDF: {pdf_path or 'không tạo được (máy không có Pages/LibreOffice) — mở file Word rồi Xuất PDF'}")
    print(f"HÌNH: {img_dir} ({images} ảnh chụp, {diagrams} sơ đồ)")
    print(f"{len(chapters)} chương, {sections} mục")
    if not all(ok for _, ok, _ in checks):
        sys.exit(2)


if __name__ == "__main__":
    main()
